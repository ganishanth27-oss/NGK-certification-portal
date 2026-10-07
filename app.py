"""NGK Portal — Flask event and certificate management application."""
from __future__ import annotations

import csv
import io
import json
import os
import re
import secrets
import sqlite3
import uuid
from datetime import date, datetime, timezone
from functools import wraps
from pathlib import Path

from dotenv import load_dotenv
from flask import Flask, jsonify, make_response, render_template, request, send_file, session
from itsdangerous import BadSignature, SignatureExpired, URLSafeTimedSerializer
from openpyxl import Workbook, load_workbook
import xlrd
from PIL import Image
from pypdf import PdfReader, PdfWriter
from reportlab.lib import colors
from reportlab.lib.pagesizes import landscape, letter
from reportlab.lib.utils import ImageReader
from reportlab.pdfgen import canvas
from werkzeug.security import check_password_hash, generate_password_hash
from werkzeug.utils import secure_filename

BASE_DIR = Path(__file__).resolve().parent
load_dotenv(BASE_DIR / ".env")


def local_path(value, default: Path) -> Path:
    path = Path(value) if value else default
    return path if path.is_absolute() else BASE_DIR / path


DATA_DIR = local_path(os.environ.get("DATA_DIR"), BASE_DIR / "data")
UPLOAD_DIR = DATA_DIR / "uploads"
DB_PATH = local_path(os.environ.get("DATABASE_PATH"), DATA_DIR / "portal.sqlite3")
DATA_DIR.mkdir(parents=True, exist_ok=True)
UPLOAD_DIR.mkdir(parents=True, exist_ok=True)
DB_PATH.parent.mkdir(parents=True, exist_ok=True)

app = Flask(__name__)
app.config.update(
    SECRET_KEY=os.environ.get("SECRET_KEY", "dev-only-change-me-please-" + secrets.token_hex(12)),
    MAX_CONTENT_LENGTH=15 * 1024 * 1024,
    SESSION_COOKIE_HTTPONLY=True,
    SESSION_COOKIE_SAMESITE="Lax",
    SESSION_COOKIE_SECURE=os.environ.get("COOKIE_SECURE", "0") == "1",
)
serializer = URLSafeTimedSerializer(app.config["SECRET_KEY"], salt="participant-certificate-v1")
ALLOWED_IMPORTS = {"xlsx", "xlsm", "xls", "csv", "txt", "pdf"}
ALLOWED_TEMPLATES = {"png", "jpg", "jpeg", "pdf"}

SCHEMA = """
PRAGMA foreign_keys = ON;
CREATE TABLE IF NOT EXISTS admins (
 id INTEGER PRIMARY KEY AUTOINCREMENT, name TEXT NOT NULL, email TEXT NOT NULL UNIQUE,
 password_hash TEXT NOT NULL, created_at TEXT NOT NULL
);
CREATE TABLE IF NOT EXISTS events (
 id INTEGER PRIMARY KEY AUTOINCREMENT, name TEXT NOT NULL, event_date TEXT NOT NULL DEFAULT '',
 venue TEXT NOT NULL DEFAULT '', description TEXT NOT NULL DEFAULT '', status TEXT NOT NULL DEFAULT 'Active',
 created_at TEXT NOT NULL
);
CREATE TABLE IF NOT EXISTS participants (
 id INTEGER PRIMARY KEY AUTOINCREMENT, event_id INTEGER NOT NULL REFERENCES events(id) ON DELETE CASCADE,
 student_name TEXT NOT NULL, mobile TEXT NOT NULL, email TEXT NOT NULL, participation TEXT NOT NULL DEFAULT '',
 certificate_id TEXT NOT NULL UNIQUE, certificate_status TEXT NOT NULL DEFAULT 'Available', created_at TEXT NOT NULL,
 UNIQUE(event_id, mobile, email)
);
CREATE TABLE IF NOT EXISTS certificate_templates (
 id INTEGER PRIMARY KEY AUTOINCREMENT, event_id INTEGER NOT NULL UNIQUE REFERENCES events(id) ON DELETE CASCADE,
 template_file TEXT, configuration TEXT NOT NULL DEFAULT '{}', created_at TEXT NOT NULL, updated_at TEXT NOT NULL
);
CREATE TABLE IF NOT EXISTS certificates (
 id INTEGER PRIMARY KEY AUTOINCREMENT, participant_id INTEGER NOT NULL UNIQUE REFERENCES participants(id) ON DELETE CASCADE,
 generated_file TEXT, generated_at TEXT, download_count INTEGER NOT NULL DEFAULT 0
);
CREATE TABLE IF NOT EXISTS settings (key TEXT PRIMARY KEY, value TEXT NOT NULL);
CREATE INDEX IF NOT EXISTS idx_participant_lookup ON participants(mobile, email);
CREATE INDEX IF NOT EXISTS idx_participant_event ON participants(event_id);
"""


def now_iso() -> str:
    return datetime.now(timezone.utc).replace(microsecond=0).isoformat()


def get_db() -> sqlite3.Connection:
    conn = sqlite3.connect(DB_PATH, timeout=20)
    conn.row_factory = sqlite3.Row
    conn.execute("PRAGMA foreign_keys = ON")
    return conn


def row_dict(row):
    return dict(row) if row is not None else None


def certificate_code() -> str:
    return f"ECP-{date.today().year}-{uuid.uuid4().hex[:8].upper()}"


def init_db() -> None:
    db = get_db()
    db.executescript(SCHEMA)
    admin_email = os.environ.get("ADMIN_EMAIL", "").strip().lower()
    admin_password = os.environ.get("ADMIN_PASSWORD", "")
    if admin_email and admin_password and not db.execute("SELECT 1 FROM admins WHERE email=?", (admin_email,)).fetchone():
        db.execute("INSERT INTO admins(name,email,password_hash,created_at) VALUES(?,?,?,?)",
                   ("Portal Administrator", admin_email, generate_password_hash(admin_password), now_iso()))
    db.commit()
    db.close()


def default_config() -> dict:
    return {"organization": "Nehru Grand Kacheri", "organizer": "Nehru Grand Kacheri",
            "title": "CERTIFICATE OF PARTICIPATION", "intro": "This certificate is proudly presented to",
            "participation_text": "For successfully participating in", "footer": "Thank you for being part of this event.",
            "name_x": 0.5, "name_y": 0.52, "event_x": 0.5, "event_y": 0.36,
            "name_size": 30, "event_size": 18, "accent": "#3b54c5"}


def admin_required(fn):
    @wraps(fn)
    def wrapped(*args, **kwargs):
        if not session.get("admin_id"):
            return jsonify(error="Please sign in to continue."), 401
        if request.method not in ("GET", "HEAD", "OPTIONS"):
            sent = request.headers.get("X-CSRF-Token", "")
            if not sent or not secrets.compare_digest(sent, session.get("csrf", "")):
                return jsonify(error="Session security token is missing or expired. Please sign in again."), 403
        return fn(*args, **kwargs)
    return wrapped


def api_error(message: str, status: int = 400):
    return jsonify(error=message), status


def normalize_mobile(value) -> str:
    return re.sub(r"\D", "", str(value or ""))


def valid_email(value: str) -> bool:
    return bool(re.fullmatch(r"[^\s@]+@[^\s@]+\.[^\s@]+", str(value or "").strip()))


def normalize_header(value: str) -> str:
    return re.sub(r"[^a-z0-9]", "", str(value or "").strip().lower())


def map_headers(headers):
    aliases = {
        "student_name": {"name", "studentname", "fullname", "participant", "participantname", "student"},
        "mobile": {"mobile", "mobilenumber", "phone", "phonenumber", "contact", "contactnumber"},
        "email": {"email", "emailid", "gmail", "gmailid", "emailaddress"},
        "participation": {"participation", "event", "eventname", "activity", "participationevent"},
    }
    result = {}
    for idx, header in enumerate(headers):
        key = normalize_header(header)
        for field, options in aliases.items():
            if key in options and field not in result:
                result[field] = idx
    return result


def rows_to_participants(rows):
    if not rows:
        return []
    headers = [str(v or "") for v in rows[0]]
    mapping = map_headers(headers)
    if "student_name" not in mapping or "mobile" not in mapping or "email" not in mapping:
        mapping = {"student_name": 0, "mobile": 1, "email": 2}
        if len(headers) > 3:
            mapping["participation"] = 3
        data_rows = rows
    else:
        data_rows = rows[1:]
    result = []
    for row in data_rows:
        vals = list(row)
        def cell(field):
            i = mapping.get(field)
            return str(vals[i] if i is not None and i < len(vals) and vals[i] is not None else "").strip()
        name, mobile, email = cell("student_name"), normalize_mobile(cell("mobile")), cell("email").lower()
        if not name and not mobile and not email:
            continue
        result.append({"student_name": name, "mobile": mobile, "email": email,
                       "participation": cell("participation")})
    return result


def parse_import(file_storage):
    original = secure_filename(file_storage.filename or "")
    ext = Path(original).suffix.lower().lstrip(".")
    if ext not in ALLOWED_IMPORTS:
        return api_error("Unsupported file. Upload XLS, XLSX, CSV, TXT or PDF.")
    raw = file_storage.read()
    if not raw:
        raise ValueError("The uploaded file is empty.")
    if len(raw) > 15 * 1024 * 1024:
        raise ValueError("File is larger than the 15 MB limit.")
    if ext == "xls":
        try:
            book = xlrd.open_workbook(file_contents=raw)
            sheet = book.sheet_by_index(0)
            rows = [[sheet.cell_value(r, c) for c in range(sheet.ncols)] for r in range(sheet.nrows)]
        except Exception as exc:
            raise ValueError("Could not read the XLS file. Please check that it is a valid workbook.") from exc
        return rows_to_participants(rows)
    if ext in {"xlsx", "xlsm"}:
        try:
            wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
            rows = list(wb.active.iter_rows(values_only=True))
            wb.close()
        except Exception as exc:
            raise ValueError("Could not read the Excel file. Please check that it is a valid workbook.") from exc
        return rows_to_participants(rows)
    if ext in {"csv", "txt"}:
        try:
            text = raw.decode("utf-8-sig")
        except UnicodeDecodeError:
            text = raw.decode("latin-1")
        try:
            dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t|")
            rows = list(csv.reader(io.StringIO(text), dialect))
        except csv.Error:
            rows = [re.split(r"\s*[|,;\t]\s*", line.strip()) for line in text.splitlines() if line.strip()]
        return rows_to_participants(rows)
    try:
        reader = PdfReader(io.BytesIO(raw))
        text = "\n".join((page.extract_text() or "") for page in reader.pages)
    except Exception as exc:
        raise ValueError("Could not read the PDF. Upload a valid text-based PDF or use CSV/XLSX.") from exc
    parsed = []
    for line in text.splitlines():
        line = line.strip()
        if not line or "@" not in line:
            continue
        email_match = re.search(r"[A-Z0-9._%+-]+@[A-Z0-9.-]+\.[A-Z]{2,}", line, re.I)
        phone_match = re.search(r"(?:\+?\d[\d\s().-]{7,}\d)", line)
        if not email_match or not phone_match:
            continue
        email = email_match.group(0)
        phone = normalize_mobile(phone_match.group(0))
        left = line[:min(email_match.start(), phone_match.start())].strip(" |,;\t-:")
        right = line[max(email_match.end(), phone_match.end()):].strip(" |,;\t-:")
        name = re.sub(r"\s{2,}", " ", left or right).strip()
        if name.lower() in {"name", "student name", "participant name"}:
            continue
        parsed.append({"student_name": name, "mobile": phone, "email": email.lower(), "participation": ""})
    return parsed


def template_config(db, event_id):
    row = db.execute("SELECT * FROM certificate_templates WHERE event_id=?", (event_id,)).fetchone()
    cfg = default_config()
    if row:
        try:
            cfg.update(json.loads(row["configuration"] or "{}"))
        except (ValueError, TypeError):
            pass
    return row, cfg


def generate_certificate(participant):
    db = get_db()
    event = db.execute("SELECT * FROM events WHERE id=?", (participant["event_id"],)).fetchone()
    template, cfg = template_config(db, participant["event_id"])
    db.close()
    if not event:
        raise ValueError("The event for this certificate no longer exists.")
    output = io.BytesIO()
    page_size = landscape(letter)
    bg_path = UPLOAD_DIR / template["template_file"] if template and template["template_file"] else None
    if bg_path and bg_path.exists() and bg_path.suffix.lower() == ".pdf":
        try:
            bg_reader = PdfReader(str(bg_path))
            page = bg_reader.pages[0]
            width = float(page.mediabox.width)
            height = float(page.mediabox.height)
        except Exception as exc:
            raise ValueError("The saved PDF template could not be opened.") from exc
    elif bg_path and bg_path.exists():
        with Image.open(bg_path) as im:
            width, height = im.size
        page_size = (width, height)
    else:
        width, height = page_size
    overlay = io.BytesIO()
    c = canvas.Canvas(overlay, pagesize=(width, height))
    c.setFillColor(colors.HexColor(cfg.get("accent", "#3b54c5")))
    c.setFont("Helvetica-Bold", 19)
    c.drawCentredString(width / 2, height * .82, str(cfg.get("title", "CERTIFICATE OF PARTICIPATION"))[:90])
    c.setFillColor(colors.HexColor("#344054"))
    c.setFont("Helvetica", 13)
    c.drawCentredString(width / 2, height * .70, str(cfg.get("intro", "This certificate is proudly presented to"))[:150])
    name_y = float(cfg.get("name_y", .52))
    name_x = float(cfg.get("name_x", .5))
    c.setFillColor(colors.HexColor(cfg.get("accent", "#3b54c5")))
    c.setFont("Helvetica-Bold", min(44, max(18, int(cfg.get("name_size", 30)))))
    c.drawCentredString(width * name_x, height * name_y, participant["student_name"][:100])
    c.setFillColor(colors.HexColor("#344054"))
    c.setFont("Helvetica", 13)
    c.drawCentredString(width / 2, height * .43, str(cfg.get("participation_text", "For successfully participating in"))[:150])
    c.setFillColor(colors.HexColor("#172554"))
    c.setFont("Helvetica-Bold", min(28, max(14, int(cfg.get("event_size", 18)))))
    c.drawCentredString(width * float(cfg.get("event_x", .5)), height * float(cfg.get("event_y", .36)), event["name"][:100])
    c.setFillColor(colors.HexColor("#475467"))
    c.setFont("Helvetica", 11)
    held = f"held on {event['event_date']}" if event["event_date"] else "in recognition of your participation"
    c.drawCentredString(width / 2, height * .27, held)
    c.setStrokeColor(colors.HexColor(cfg.get("accent", "#3b54c5")))
    c.line(width * .34, height * .19, width * .66, height * .19)
    c.setFont("Helvetica-Bold", 10)
    c.drawCentredString(width / 2, height * .155, str(cfg.get("organizer", "Event Organizing Committee"))[:100])
    c.setFont("Helvetica", 9)
    c.drawCentredString(width / 2, height * .105, str(cfg.get("organization", "Event Certificate Portal"))[:100])
    c.setFillColor(colors.HexColor("#667085"))
    c.setFont("Helvetica", 8)
    c.drawString(26, 20, f"Certificate ID: {participant['certificate_id']}")
    if cfg.get("footer"):
        c.drawRightString(width - 26, 20, str(cfg["footer"])[:100])
    c.save()
    overlay.seek(0)
    over_reader = PdfReader(overlay)
    page = over_reader.pages[0]
    if bg_path and bg_path.exists() and bg_path.suffix.lower() == ".pdf":
        bg = PdfReader(str(bg_path)).pages[0]
        bg.merge_page(page)
        page = bg
    elif bg_path and bg_path.exists():
        bg_pdf = io.BytesIO()
        bc = canvas.Canvas(bg_pdf, pagesize=(width, height))
        bc.drawImage(ImageReader(str(bg_path)), 0, 0, width=width, height=height, preserveAspectRatio=False)
        bc.save()
        bg_page = PdfReader(bg_pdf).pages[0]
        bg_page.merge_page(page)
        page = bg_page
    writer = PdfWriter()
    writer.add_page(page)
    writer.write(output)
    output.seek(0)
    return output


@app.get("/")
def index():
    return render_template("index.html")


@app.get("/admin")
def admin_page():
    return render_template("admin.html")


@app.get("/api/session")
def session_state():
    return jsonify(authenticated=bool(session.get("admin_id")), email=session.get("admin_email"), csrf=session.get("csrf"))


@app.post("/api/login")
def login():
    body = request.get_json(silent=True) or {}
    email = str(body.get("email", "")).strip().lower()
    password = str(body.get("password", ""))
    db = get_db()
    admin = db.execute("SELECT * FROM admins WHERE email=?", (email,)).fetchone()
    db.close()
    if not admin or not check_password_hash(admin["password_hash"], password):
        return api_error("Email or password is incorrect.", 401)
    session.clear()
    session["admin_id"] = admin["id"]
    session["admin_email"] = admin["email"]
    session["csrf"] = secrets.token_urlsafe(32)
    return jsonify(authenticated=True, email=admin["email"], csrf=session["csrf"])


@app.post("/api/logout")
@admin_required
def logout():
    session.clear()
    return jsonify(ok=True)


@app.get("/api/admin/summary")
@admin_required
def summary():
    db = get_db()
    counts = db.execute("SELECT (SELECT COUNT(*) FROM events) events,(SELECT COUNT(*) FROM participants) participants,(SELECT COUNT(*) FROM participants) certificates,(SELECT COALESCE(SUM(download_count),0) FROM certificates) downloads").fetchone()
    recent_events = db.execute("SELECT e.*,COUNT(p.id) participant_count FROM events e LEFT JOIN participants p ON p.event_id=e.id GROUP BY e.id ORDER BY e.created_at DESC LIMIT 5").fetchall()
    recent_people = db.execute("SELECT p.*,e.name event_name FROM participants p JOIN events e ON e.id=p.event_id ORDER BY p.created_at DESC LIMIT 6").fetchall()
    db.close()
    return jsonify(stats=dict(counts), recent_events=[row_dict(r) for r in recent_events], recent_participants=[row_dict(r) for r in recent_people])


@app.get("/api/admin/events")
@admin_required
def list_events():
    db = get_db()
    rows = db.execute("SELECT e.*,COUNT(p.id) participant_count FROM events e LEFT JOIN participants p ON p.event_id=e.id GROUP BY e.id ORDER BY e.created_at DESC").fetchall()
    db.close()
    return jsonify(events=[row_dict(r) for r in rows])


@app.post("/api/admin/events")
@admin_required
def create_event():
    body = request.get_json(silent=True) or {}
    name = str(body.get("name", "")).strip()
    if not name:
        return api_error("Event name is required.")
    db = get_db()
    cur = db.execute("INSERT INTO events(name,event_date,venue,description,status,created_at) VALUES(?,?,?,?,?,?)",
                     (name, str(body.get("event_date", ""))[:20], str(body.get("venue", ""))[:200],
                      str(body.get("description", ""))[:2000], str(body.get("status", "Active"))[:30], now_iso()))
    eid = cur.lastrowid
    db.execute("INSERT INTO certificate_templates(event_id,configuration,created_at,updated_at) VALUES(?,?,?,?)",
               (eid, json.dumps(default_config()), now_iso(), now_iso()))
    db.commit()
    row = db.execute("SELECT * FROM events WHERE id=?", (eid,)).fetchone()
    db.close()
    return jsonify(event=row_dict(row)), 201


@app.put("/api/admin/events/<int:event_id>")
@admin_required
def update_event(event_id):
    body = request.get_json(silent=True) or {}
    db = get_db()
    if not db.execute("SELECT 1 FROM events WHERE id=?", (event_id,)).fetchone():
        db.close()
        return api_error("Event not found.", 404)
    fields = ["name", "event_date", "venue", "description", "status"]
    current = row_dict(db.execute("SELECT * FROM events WHERE id=?", (event_id,)).fetchone())
    values = {key: str(body.get(key, current[key])).strip() for key in fields}
    if not values["name"]:
        db.close()
        return api_error("Event name is required.")
    db.execute("UPDATE events SET name=?,event_date=?,venue=?,description=?,status=? WHERE id=?",
               (values["name"], values["event_date"][:20], values["venue"][:200], values["description"][:2000], values["status"][:30], event_id))
    db.commit()
    row = row_dict(db.execute("SELECT * FROM events WHERE id=?", (event_id,)).fetchone())
    db.close()
    return jsonify(event=row)


@app.delete("/api/admin/events/<int:event_id>")
@admin_required
def delete_event(event_id):
    db = get_db()
    cur = db.execute("DELETE FROM events WHERE id=?", (event_id,))
    db.commit()
    db.close()
    return (jsonify(ok=True) if cur.rowcount else api_error("Event not found.", 404))


@app.get("/api/admin/participants")
@admin_required
def list_participants():
    search = str(request.args.get("q", "")).strip()
    event_id = request.args.get("event_id", "")
    sql = "SELECT p.*,e.name event_name FROM participants p JOIN events e ON e.id=p.event_id WHERE 1=1"
    args = []
    if search:
        needle = f"%{search}%"
        sql += " AND (p.student_name LIKE ? OR p.mobile LIKE ? OR p.email LIKE ?)"
        args.extend([needle] * 3)
    if event_id.isdigit():
        sql += " AND p.event_id=?"
        args.append(int(event_id))
    sql += " ORDER BY p.created_at DESC"
    db = get_db()
    rows = db.execute(sql, args).fetchall()
    db.close()
    return jsonify(participants=[row_dict(r) for r in rows])


@app.post("/api/admin/import/preview")
@admin_required
def import_preview():
    file = request.files.get("file")
    if not file:
        return api_error("Choose a participant file to upload.")
    try:
        rows = parse_import(file)
    except ValueError as exc:
        return api_error(str(exc))
    if not rows:
        return api_error("No participant rows could be detected. Use columns: Name, Mobile, Email, Participation.")
    if len(rows) > 10000:
        return api_error("This file has more than 10,000 rows. Split it into smaller files.")
    return jsonify(filename=secure_filename(file.filename or "participants"), count=len(rows), rows=rows)


@app.post("/api/admin/import/commit")
@admin_required
def import_commit():
    body = request.get_json(silent=True) or {}
    event_id = body.get("event_id")
    rows = body.get("rows")
    if not isinstance(rows, list) or not rows:
        return api_error("There are no participant rows to import.")
    if len(rows) > 10000:
        return api_error("A maximum of 10,000 rows can be imported at once.")
    db = get_db()
    if not db.execute("SELECT 1 FROM events WHERE id=?", (event_id,)).fetchone():
        db.close()
        return api_error("Select a valid event.")
    imported, skipped, errors = 0, 0, []
    for idx, row in enumerate(rows, start=1):
        if not isinstance(row, dict):
            errors.append(f"Row {idx}: invalid row data")
            continue
        name = str(row.get("student_name", "")).strip()[:200]
        mobile = normalize_mobile(row.get("mobile"))
        email = str(row.get("email", "")).strip().lower()[:254]
        participation = str(row.get("participation", "")).strip()[:200]
        if not name or len(mobile) < 7 or not valid_email(email):
            errors.append(f"Row {idx}: name, valid mobile and email are required")
            continue
        try:
            cur = db.execute("INSERT INTO participants(event_id,student_name,mobile,email,participation,certificate_id,created_at) VALUES(?,?,?,?,?,?,?)",
                             (int(event_id), name, mobile, email, participation, certificate_code(), now_iso()))
            db.execute("INSERT INTO certificates(participant_id) VALUES(?)", (cur.lastrowid,))
            imported += 1
        except sqlite3.IntegrityError:
            skipped += 1
    db.commit()
    db.close()
    return jsonify(imported=imported, skipped=skipped, errors=errors[:50])


@app.put("/api/admin/participants/<int:participant_id>")
@admin_required
def update_participant(participant_id):
    body = request.get_json(silent=True) or {}
    name = str(body.get("student_name", "")).strip()[:200]
    mobile = normalize_mobile(body.get("mobile"))
    email = str(body.get("email", "")).strip().lower()[:254]
    participation = str(body.get("participation", "")).strip()[:200]
    event_id = body.get("event_id")
    if not name or len(mobile) < 7 or not valid_email(email):
        return api_error("Enter a name, valid mobile number and email address.")
    db = get_db()
    try:
        cur = db.execute("UPDATE participants SET student_name=?,mobile=?,email=?,participation=?,event_id=? WHERE id=?",
                         (name, mobile, email, participation, int(event_id), participant_id))
        db.commit()
    except (sqlite3.IntegrityError, TypeError, ValueError):
        db.close()
        return api_error("Could not save participant. Check the event and duplicate records.")
    db.close()
    return jsonify(ok=bool(cur.rowcount)) if cur.rowcount else api_error("Participant not found.", 404)


@app.delete("/api/admin/participants/<int:participant_id>")
@admin_required
def delete_participant(participant_id):
    db = get_db()
    cur = db.execute("DELETE FROM participants WHERE id=?", (participant_id,))
    db.commit()
    db.close()
    return jsonify(ok=True) if cur.rowcount else api_error("Participant not found.", 404)


@app.get("/api/admin/export")
@admin_required
def export_participants():
    fmt = request.args.get("format", "csv").lower()
    db = get_db()
    rows = db.execute("SELECT p.student_name,p.mobile,p.email,e.name event,p.participation,p.certificate_status,p.certificate_id FROM participants p JOIN events e ON e.id=p.event_id ORDER BY e.name,p.student_name").fetchall()
    db.close()
    headers = ["Student Name", "Mobile Number", "Email", "Event", "Participation", "Certificate Status", "Certificate ID"]
    if fmt == "xlsx":
        wb = Workbook()
        ws = wb.active
        ws.title = "Participants"
        ws.append(headers)
        for row in rows:
            ws.append(list(row))
        out = io.BytesIO()
        wb.save(out)
        out.seek(0)
        return send_file(out, as_attachment=True, download_name="participants.xlsx", mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    out = io.StringIO()
    writer = csv.writer(out)
    writer.writerow(headers)
    writer.writerows([list(r) for r in rows])
    response = make_response("\ufeff" + out.getvalue())
    response.headers["Content-Type"] = "text/csv; charset=utf-8"
    response.headers["Content-Disposition"] = "attachment; filename=participants.csv"
    return response


@app.get("/api/admin/templates/<int:event_id>")
@admin_required
def get_template(event_id):
    db = get_db()
    if not db.execute("SELECT 1 FROM events WHERE id=?", (event_id,)).fetchone():
        db.close()
        return api_error("Event not found.", 404)
    row, cfg = template_config(db, event_id)
    db.close()
    return jsonify(event_id=event_id, uploaded=bool(row and row["template_file"]), filename=Path(row["template_file"]).name if row and row["template_file"] else None, configuration=cfg)


@app.post("/api/admin/templates/<int:event_id>")
@admin_required
def save_template(event_id):
    db = get_db()
    if not db.execute("SELECT 1 FROM events WHERE id=?", (event_id,)).fetchone():
        db.close()
        return api_error("Event not found.", 404)
    old, cfg = template_config(db, event_id)
    uploaded = request.files.get("file")
    file_name = old["template_file"] if old else None
    if uploaded and uploaded.filename:
        safe = secure_filename(uploaded.filename)
        ext = Path(safe).suffix.lower().lstrip(".")
        if ext not in ALLOWED_TEMPLATES:
            db.close()
            return api_error("Template must be a PNG, JPG or PDF file.")
        data = uploaded.read()
        if not data or len(data) > 12 * 1024 * 1024:
            db.close()
            return api_error("Template file is empty or larger than 12 MB.")
        file_name = f"template-{event_id}-{uuid.uuid4().hex[:12]}.{ext}"
        target = UPLOAD_DIR / file_name
        try:
            target.write_bytes(data)
            if ext in {"png", "jpg", "jpeg"}:
                with Image.open(target) as im:
                    im.verify()
            else:
                PdfReader(str(target))
        except Exception:
            target.unlink(missing_ok=True)
            db.close()
            return api_error("That template file is invalid or could not be read.")
    body = request.form
    for key in ("organization", "organizer", "title", "intro", "participation_text", "footer"):
        if body.get(key) is not None:
            cfg[key] = body.get(key, "")[:250]
    for key, lo, hi in (("name_size", 18, 44), ("event_size", 14, 28)):
        try:
            cfg[key] = max(lo, min(hi, int(body.get(key, cfg.get(key, lo)))))
        except (TypeError, ValueError):
            pass
    accent = body.get("accent", cfg.get("accent", "#3b54c5"))
    if re.fullmatch(r"#[0-9a-fA-F]{6}", accent):
        cfg["accent"] = accent
    if old:
        db.execute("UPDATE certificate_templates SET template_file=?,configuration=?,updated_at=? WHERE event_id=?",
                   (file_name, json.dumps(cfg), now_iso(), event_id))
    else:
        db.execute("INSERT INTO certificate_templates(event_id,template_file,configuration,created_at,updated_at) VALUES(?,?,?,?,?)",
                   (event_id, file_name, json.dumps(cfg), now_iso(), now_iso()))
    db.commit()
    db.close()
    return jsonify(ok=True, uploaded=bool(file_name), filename=Path(file_name).name if file_name else None, configuration=cfg)


@app.post("/api/verify")
def verify_participant():
    body = request.get_json(silent=True) or {}
    mobile = normalize_mobile(body.get("mobile"))
    email = str(body.get("email", "")).strip().lower()
    if len(mobile) < 7 or not valid_email(email):
        return api_error("Enter a valid mobile number and email address.")
    db = get_db()
    rows = db.execute("SELECT p.id,p.student_name,p.event_id,p.participation,p.certificate_id,p.certificate_status,e.name event_name,e.event_date,e.venue FROM participants p JOIN events e ON e.id=p.event_id WHERE p.mobile=? AND lower(p.email)=? ORDER BY e.event_date DESC,p.id DESC", (mobile, email)).fetchall()
    db.close()
    # Same response regardless of whether one or both identifiers were present.
    if not rows:
        return api_error("Certificate not found. Please check your mobile number and email address and try again.", 404)
    found = []
    for r in rows:
        token = serializer.dumps({"participant_id": r["id"]})
        found.append({"student_name": r["student_name"], "event_name": r["event_name"], "event_date": r["event_date"],
                      "participation": r["participation"], "certificate_id": r["certificate_id"],
                      "certificate_status": r["certificate_status"], "token": token,
                      "view_url": f"/certificate/{token}", "download_url": f"/api/certificate/{token}.pdf"})
    return jsonify(student_name=rows[0]["student_name"], certificates=found)


def participant_from_token(token):
    try:
        data = serializer.loads(token, max_age=600)
    except SignatureExpired:
        return None, "Verification expired. Please verify your mobile and email again."
    except BadSignature:
        return None, "Certificate link is invalid. Please verify your mobile and email again."
    db = get_db()
    row = db.execute("SELECT p.*,e.name event_name,e.event_date FROM participants p JOIN events e ON e.id=p.event_id WHERE p.id=?", (data.get("participant_id"),)).fetchone()
    db.close()
    return row, None


@app.get("/certificate/<token>")
def certificate_preview(token):
    row, error = participant_from_token(token)
    if error or not row:
        return render_template("certificate_error.html", message=error or "Certificate not found."), 404
    db = get_db()
    _, cfg = template_config(db, row["event_id"])
    db.close()
    return render_template("certificate.html", participant=row_dict(row), token=token, config=cfg)


@app.get("/api/certificate/<token>.pdf")
def certificate_pdf(token):
    row, error = participant_from_token(token)
    if error or not row:
        return api_error(error or "Certificate not found.", 404)
    try:
        pdf = generate_certificate(row)
    except Exception as exc:
        app.logger.exception("Certificate generation failed")
        return api_error(str(exc) if isinstance(exc, ValueError) else "Certificate could not be generated.", 500)
    db = get_db()
    db.execute("INSERT INTO certificates(participant_id,generated_at,download_count) VALUES(?,?,1) ON CONFLICT(participant_id) DO UPDATE SET generated_at=excluded.generated_at,download_count=download_count+1", (row["id"], now_iso()))
    db.commit()
    db.close()
    file_stub = re.sub(r"[^A-Za-z0-9_-]+", "_", f"{row['student_name']}_{row['event_name']}").strip("_")[:90]
    return send_file(pdf, as_attachment=request.args.get("download") == "1", download_name=f"{file_stub}_Certificate.pdf", mimetype="application/pdf")


@app.errorhandler(413)
def too_large(_error):
    return api_error("Uploaded file is too large. The maximum size is 15 MB.", 413)


@app.errorhandler(500)
def server_error(_error):
    return jsonify(error="Unexpected server error. Please try again."), 500


init_db()
if __name__ == "__main__":
    app.run(host=os.environ.get("HOST", "127.0.0.1"), port=int(os.environ.get("PORT", 5000)), debug=os.environ.get("FLASK_DEBUG", "0") == "1")
